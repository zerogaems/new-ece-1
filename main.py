import io
import json
import os
import re
import time
import zipfile
import logging
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from threading import Thread

import libsql
import telebot
from flask import Flask
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from telebot import types

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
)
logger = logging.getLogger('freshmen_bot')

# ==================== خادم Flask لإرضاء Render و UptimeRobot ====================
app = Flask('')


@app.route('/')
def home():
  return 'Freshman Verification Bot - Fully Secure & Active!'


@app.route('/healthz')
def healthz():
  return 'ok'


def run_flask():
  port = int(os.environ.get('PORT', 8080))
  app.run(host='0.0.0.0', port=port)


# ==================== الإعدادات الأساسية (متغيرات البيئة) ====================
BOT_TOKEN = os.environ.get('BOT_TOKEN', '').strip()
if not BOT_TOKEN:
  raise RuntimeError('يجب ضبط متغير البيئة BOT_TOKEN قبل تشغيل البوت.')

# الأدمنية الأساسيون (Super Admins): آيديات مفصولة بفواصل.
# هدول بيقدروا يضيفوا/يحذفوا أدمنية من داخل البوت.
SUPER_ADMIN_IDS = {
    int(x.strip())
    for x in os.environ.get('ADMIN_IDS', '7547218555').split(',')
    if x.strip().isdigit()
}

# قنوات/مجموعات البداية: بتنزرع بالقاعدة أول مرة بس، وبعدها بتتحكم فيها من
# (⚙️ الإعدادات ← القنوات والمجموعات) بدون ما تلمس الكود.
FRESHMAN_LECTURES_ID = int(
    os.environ.get('FRESHMAN_LECTURES_ID', '-1004413316628')
)
FRESHMAN_DISCUSSION_ID = int(
    os.environ.get('FRESHMAN_DISCUSSION_ID', '-1003953300954')
)

DEFAULT_BRANCH_NAME = os.environ.get(
    'BRANCH_NAME', 'الهندسة الإلكترونية والاتصالات'
)
DEFAULT_OWNER_USERNAME = os.environ.get('OWNER_USERNAME', 'Youssef_Sabra')

MAX_ATTEMPTS = int(os.environ.get('MAX_ATTEMPTS', '5'))  # محاولات التقديم
MAX_LINK_REGEN = int(os.environ.get('MAX_LINK_REGEN', '1'))  # كم مرة الطالب يطلب روابط جديدة (0 = ممنوع)
INVITE_TTL_HOURS = int(os.environ.get('INVITE_TTL_HOURS', '72'))  # صلاحية الرابط
PAGE_SIZE = 8
ZIP_PART_LIMIT = 45 * 1024 * 1024  # حد تيليجرام للملف 50MB
LOCAL_TZ = timezone(timedelta(hours=int(os.environ.get('TZ_OFFSET_HOURS', '3'))))

# تأثير الاحتفال 🎉 (message effect) عند القبول
EFFECT_CONFETTI = '5046509860389126442'

bot = telebot.TeleBot(BOT_TOKEN, num_threads=4)
admin_input_states = {}  # {admin_id: {'state': str, 'data': dict}}

# ==================== الاتصال بقاعدة بيانات Turso ====================
TURSO_URL = os.environ.get('TURSO_DATABASE_URL', '').strip()
TURSO_AUTH_TOKEN = os.environ.get('TURSO_AUTH_TOKEN', '').strip()
if not TURSO_URL or not TURSO_AUTH_TOKEN:
  raise RuntimeError(
      'يجب ضبط TURSO_DATABASE_URL و TURSO_AUTH_TOKEN كمتغيرات بيئة قبل'
      ' التشغيل.'
  )


def db_connect():
  """يفتح اتصال جديد بقاعدة Turso مع إعادة محاولة لو صار خطأ عابر."""
  last_err = None
  for attempt in range(3):
    try:
      return libsql.connect(database=TURSO_URL, auth_token=TURSO_AUTH_TOKEN)
    except Exception as e:
      last_err = e
      logger.warning('db_connect attempt %s failed: %s', attempt + 1, e)
      time.sleep(0.5 * (attempt + 1))
  raise last_err


@contextmanager
def db():
  """اتصال بينسكر لحاله حتى لو صار خطأ (ما في تسريب اتصالات)."""
  conn = db_connect()
  try:
    yield conn
  finally:
    try:
      conn.close()
    except Exception:
      pass


def q_all(sql, params=()):
  for attempt in range(2):
    try:
      with db() as conn:
        cur = conn.cursor()
        cur.execute(sql, params)
        return cur.fetchall()
    except Exception:
      if attempt == 1:
        raise
      time.sleep(0.4)


def q_one(sql, params=()):
  rows = q_all(sql, params)
  return rows[0] if rows else None


def q_exec(sql, params=()):
  """ينفذ أمر كتابة ويرجع عدد الصفوف المتأثرة."""
  with db() as conn:
    cur = conn.cursor()
    cur.execute(sql, params)
    conn.commit()
    return cur.rowcount


# ==================== أدوات الوقت والنصوص ====================
def now_utc_str():
  return datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')


def to_local(ts):
  """يحول وقت UTC المخزن لتوقيت سوريا بصيغة مقروءة."""
  if not ts:
    return ''
  try:
    dt = datetime.strptime(str(ts)[:19], '%Y-%m-%d %H:%M:%S')
    return dt.replace(tzinfo=timezone.utc).astimezone(LOCAL_TZ).strftime(
        '%Y-%m-%d %H:%M'
    )
  except Exception:
    return str(ts)


def clean_phone(phone_str):
  if not phone_str:
    return ''
  digits = re.sub(r'\D', '', str(phone_str))
  if digits.startswith('00963'):
    digits = '0' + digits[5:]
  elif digits.startswith('963'):
    digits = '0' + digits[3:]
  if len(digits) == 9 and digits.startswith('9'):
    digits = '0' + digits
  return digits


def display_name(user):
  parts = [user.first_name or '', user.last_name or '']
  name = ' '.join(p for p in parts if p).strip()
  return name or (f'@{user.username}' if user.username else str(user.id))


def safe_filename(text):
  return re.sub(r'[\\/:*?"<>|\s]+', '_', text or '').strip('_')[:60] or 'student'


STATUS_ICONS = {
    'approved': '✅',
    'pending': '⏳',
    'rejected': '❌',
    'banned': '🚫',
    'reset': '🔓',
}
STATUS_LABELS = {
    'approved': 'مقبول',
    'pending': 'قيد المراجعة',
    'rejected': 'مرفوض',
    'banned': 'محظور',
    'reset': 'مفتوح للتسجيل من جديد',
}


def status_icon(s):
  return STATUS_ICONS.get(s, '❔')


def status_label(s):
  return f'{status_icon(s)} {STATUS_LABELS.get(s, s)}'


# ==================== الإيموجي المميز (Premium) + الأزرار الملونة ====================
# الفكرة: الأدمن بيبعت للبوت رسالة فيها إيموجيات مميزة، والبوت بيربط كل إيموجي
# مميز بالإيموجي العادي المقابل له. بعدها أي إيموجي عادي بنصوص وأزرار البوت
# بيظهر مميز تلقائياً (وبالأزرار بيصير أيقونة الزر). لازم صاحب البوت يكون Premium.
# الأزرار الملونة (style) مو مرتبطة بالـ Premium، بتشتغل للكل.
BOT_EMOJIS = (
    '🎉 ✅ ❌ ⏳ 🎓 📱 🪪 📄 🔗 📚 💬 📊 📁 📥 📋 🗂 ⚙️ ⚠️ 🚫 🔍 📣 🔓 🛠 📜 🚀 👋 🔙 '
    '👥 📡 🎭 🏷 👤 🕒 📌 🔁 ♻️ 🔄 ➕ 🗑 ✏️ 📝 📷 ⛔ 👍 😢 ⭐ 📎 🗓 ⏰ 📭 📦 👆 👇 ◀️ ▶️'
)
PREMIUM_BROKEN = {'on': False, 'reason': ''}
_pe_cache = {'raw': None, 'map': {}, 'rx': None}


def u16len(text):
  return len(text.encode('utf-16-le')) // 2


def norm_emoji(ch):
  return ch.replace('\ufe0f', '')


def pe_data():
  raw = get_setting('pe_map', '')
  if raw != _pe_cache['raw']:
    try:
      mp = json.loads(raw) if raw else {}
    except Exception:
      mp = {}
    keys = sorted(mp.keys(), key=len, reverse=True)
    rx = (re.compile('|'.join(re.escape(k) + '\ufe0f?' for k in keys))
          if keys else None)
    _pe_cache.update(raw=raw, map=mp, rx=rx)
  return _pe_cache['map'], _pe_cache['rx']


def premium_active():
  return (get_setting('premium_mode', 'on') == 'on'
          and not PREMIUM_BROKEN['on'] and bool(pe_data()[0]))


def rich(text):
  """يرجع (النص، entities) مع تحويل الإيموجي العادي المربوط لإيموجي مميز."""
  if not text or not premium_active():
    return text, None
  mp, rx = pe_data()
  ents, out, pos, last = [], [], 0, 0
  for mt in rx.finditer(text):
    seg = text[last:mt.start()]
    out.append(seg)
    pos += u16len(seg)
    tok = mt.group(0)
    out.append(tok)
    ents.append(types.MessageEntity(
        type='custom_emoji', offset=pos, length=u16len(tok),
        custom_emoji_id=mp[norm_emoji(tok)],
    ))
    pos += u16len(tok)
    last = mt.end()
  out.append(text[last:])
  return ''.join(out), (ents[:100] or None)


def _icon_for(text):
  """لو النص بيبدأ بإيموجي مربوط: يرجع (icon_id, النص بدون الإيموجي)."""
  if not premium_active():
    return None, text
  mp, rx = pe_data()
  mt = rx.match(text or '')
  if not mt:
    return None, text
  label = text[mt.end():].lstrip()
  if not label:
    return None, text
  return mp[norm_emoji(mt.group(0))], label


def _attempts(text, style):
  """محاولات بناء الزر من الأغنى للأبسط (حسب دعم نسخة المكتبة)."""
  icon_id, label = _icon_for(text)
  attempts = []
  if icon_id:
    extra = {'icon_custom_emoji_id': icon_id}
    if style:
      extra['style'] = style
    attempts.append((label, extra))
  if style:
    attempts.append((text, {'style': style}))
  attempts.append((text, {}))
  return attempts


def btn(text, callback_data=None, url=None, style=None):
  """زر inline. style: 'primary' (أزرق) | 'success' (أخضر) | 'danger' (أحمر).
  والإيموجي اللي بأول النص بيصير أيقونة مميزة للزر إذا كان مربوط."""
  kwargs = {}
  if callback_data is not None:
    kwargs['callback_data'] = callback_data
  if url is not None:
    kwargs['url'] = url
  for t, extra in _attempts(text, style):
    try:
      b = types.InlineKeyboardButton(text=t, **kwargs, **extra)
      if 'icon_custom_emoji_id' in extra:
        b._fallback_text = text
      return b
    except TypeError:
      continue


def kbtn(text, request_contact=False, style=None):
  for t, extra in _attempts(text, style):
    try:
      b = types.KeyboardButton(text=t, request_contact=request_contact,
                               **extra)
      if 'icon_custom_emoji_id' in extra:
        b._fallback_text = text
      return b
    except TypeError:
      continue


def _buttons(markup):
  rows = getattr(markup, 'keyboard', None) or getattr(markup, 'rows', None)
  for row in rows or []:
    for b in row:
      yield b


def _has_icons(markup):
  return markup is not None and any(
      getattr(b, 'icon_custom_emoji_id', None) for b in _buttons(markup)
  )


def _strip_icons(markup):
  for b in _buttons(markup):
    if getattr(b, 'icon_custom_emoji_id', None):
      b.icon_custom_emoji_id = None
      b.text = getattr(b, '_fallback_text', b.text)


def _is_premium_error(e):
  low = str(e).lower()
  return any(k in low for k in ('custom emoji', 'custom_emoji', 'entit',
                                'icon', 'emoji', 'button_'))


def _premium_guard(do, ents, markup):
  """do(plain) ينفذ الإرسال. لو تيليجرام رفض الإيموجي المميز، بيرجع لنص عادي."""
  try:
    return do(False)
  except Exception as e:
    if (ents or _has_icons(markup)) and _is_premium_error(e):
      PREMIUM_BROKEN.update(on=True, reason=str(e)[:200])
      logger.warning('Premium emoji rejected by Telegram, disabled: %s', e)
      if markup is not None:
        _strip_icons(markup)
      return do(True)
    raise


def send_msg(chat_id, text, **kwargs):
  t, ents = rich(text)
  return _premium_guard(
      lambda plain: bot.send_message(
          chat_id, text if plain else t,
          **({} if plain or not ents else {'entities': ents}), **kwargs),
      ents, kwargs.get('reply_markup'),
  )


def reply_msg(message, text, **kwargs):
  return send_msg(message.chat.id, text, **kwargs)


def send_photo_rich(chat_id, photo, caption=None, **kwargs):
  t, ents = rich(caption) if caption else (caption, None)
  return _premium_guard(
      lambda plain: bot.send_photo(
          chat_id, photo, caption=caption if plain else t,
          **({} if plain or not ents else {'caption_entities': ents}),
          **kwargs),
      ents, kwargs.get('reply_markup'),
  )


# ==================== تهيئة قاعدة البيانات ====================
def ensure_columns(table, cols):
  existing = {r[1] for r in q_all(f'PRAGMA table_info({table})')}
  for name, ddl in cols.items():
    if name not in existing:
      q_exec(f'ALTER TABLE {table} ADD COLUMN {name} {ddl}')


def init_db():
  q_exec("""
      CREATE TABLE IF NOT EXISTS freshmen (
          telegram_id INTEGER PRIMARY KEY,
          full_name TEXT,
          phone TEXT,
          admission_photo_id TEXT,
          id_photo_id TEXT,
          status TEXT DEFAULT 'pending',
          created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
      )
  """)
  # ترقية الجدول القديم بدون ما نخسر أي بيانات
  ensure_columns('freshmen', {
      'admission_photo_id': 'TEXT',
      'id_photo_id': 'TEXT',
      'username': 'TEXT',
      'attempts': 'INTEGER DEFAULT 0',
      'links_count': 'INTEGER DEFAULT 0',
      'approved_by': 'TEXT',
      'approved_at': 'TEXT',
      'reject_reason': 'TEXT',
      'updated_at': 'TEXT',
  })
  q_exec("""
      CREATE TABLE IF NOT EXISTS sessions (
          telegram_id INTEGER PRIMARY KEY,
          step TEXT,
          name TEXT,
          phone TEXT,
          admission_photo_id TEXT,
          updated_at TEXT
      )
  """)
  q_exec("""
      CREATE TABLE IF NOT EXISTS channels (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          chat_id INTEGER UNIQUE,
          title TEXT,
          kind TEXT DEFAULT 'channel'
      )
  """)
  q_exec("""
      CREATE TABLE IF NOT EXISTS settings (
          key TEXT PRIMARY KEY,
          value TEXT
      )
  """)
  q_exec("""
      CREATE TABLE IF NOT EXISTS admins (
          telegram_id INTEGER PRIMARY KEY,
          name TEXT,
          added_by INTEGER,
          added_at TEXT
      )
  """)
  q_exec("""
      CREATE TABLE IF NOT EXISTS audit_log (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          ts TEXT,
          actor_id INTEGER,
          actor_name TEXT,
          action TEXT,
          target_id INTEGER,
          details TEXT
      )
  """)
  q_exec("""
      CREATE TABLE IF NOT EXISTS submissions (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          telegram_id INTEGER,
          full_name TEXT,
          phone TEXT,
          admission_photo_id TEXT,
          id_photo_id TEXT,
          submitted_at TEXT,
          result TEXT DEFAULT 'pending',
          reason TEXT
      )
  """)
  q_exec("""
      CREATE TABLE IF NOT EXISTS invite_links (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          telegram_id INTEGER,
          chat_id INTEGER,
          link TEXT,
          created_at TEXT,
          revoked INTEGER DEFAULT 0
      )
  """)
  # نزرع القناتين الأساسيتين أول مرة بس
  row = q_one('SELECT COUNT(*) FROM channels')
  if row and row[0] == 0:
    q_exec(
        'INSERT INTO channels (chat_id, title, kind) VALUES (?, ?, ?)',
        (FRESHMAN_LECTURES_ID, 'قناة المحاضرات (سنة أولى)', 'channel'),
    )
    q_exec(
        'INSERT INTO channels (chat_id, title, kind) VALUES (?, ?, ?)',
        (FRESHMAN_DISCUSSION_ID, 'مجموعة المناقشة (سنة أولى)', 'group'),
    )
  q_exec(
      'DELETE FROM sessions WHERE updated_at < ?',
      ((datetime.now(timezone.utc) - timedelta(days=7)).strftime(
          '%Y-%m-%d %H:%M:%S'
      ),),
  )


init_db()

# ==================== كاش الأدمنية والإعدادات والجلسات ====================
_admin_ids = set(SUPER_ADMIN_IDS)


def reload_admins():
  ids = set(SUPER_ADMIN_IDS)
  for (tid,) in q_all('SELECT telegram_id FROM admins'):
    ids.add(int(tid))
  _admin_ids.clear()
  _admin_ids.update(ids)


def is_admin(user_id):
  return user_id in _admin_ids


def is_super_admin(user_id):
  return user_id in SUPER_ADMIN_IDS


def notify_all_admins(send_func):
  """يستدعي send_func(admin_id) لكل أدمن ويتجاوز أي أدمن فشل الإرسال له."""
  for admin_id in list(_admin_ids):
    try:
      send_func(admin_id)
    except Exception as e:
      logger.warning('فشل إشعار الأدمن %s: %s', admin_id, e)


_settings = {}


def load_settings():
  _settings.clear()
  for k, v in q_all('SELECT key, value FROM settings'):
    _settings[k] = v


def get_setting(key, default=''):
  return _settings.get(key, default)


def set_setting(key, value):
  q_exec(
      'INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO'
      ' UPDATE SET value = excluded.value',
      (key, value),
  )
  _settings[key] = value


def branch_name():
  return get_setting('branch_name', DEFAULT_BRANCH_NAME)


def owner_tag():
  name = get_setting('owner_username', DEFAULT_OWNER_USERNAME).lstrip('@')
  return f'@{name}' if name else ''


sessions = {}  # كاش بالذاكرة، ومحفوظ بالقاعدة عشان ما يضيع عند إعادة التشغيل


def load_sessions():
  sessions.clear()
  for tid, step, name, phone, adm in q_all(
      'SELECT telegram_id, step, name, phone, admission_photo_id FROM sessions'
  ):
    sessions[int(tid)] = {
        'step': step, 'name': name, 'phone': phone, 'admission_photo': adm
    }


def session_step(user_id):
  s = sessions.get(user_id)
  return s['step'] if s else None


def save_session(user_id, **fields):
  s = sessions.setdefault(user_id, {'step': 'NAME'})
  s.update(fields)
  q_exec(
      'INSERT INTO sessions (telegram_id, step, name, phone,'
      ' admission_photo_id, updated_at) VALUES (?, ?, ?, ?, ?, ?) ON'
      ' CONFLICT(telegram_id) DO UPDATE SET step = excluded.step, name ='
      ' excluded.name, phone = excluded.phone, admission_photo_id ='
      ' excluded.admission_photo_id, updated_at = excluded.updated_at',
      (user_id, s.get('step'), s.get('name'), s.get('phone'),
       s.get('admission_photo'), now_utc_str()),
  )


def drop_session(user_id):
  sessions.pop(user_id, None)
  try:
    q_exec('DELETE FROM sessions WHERE telegram_id = ?', (user_id,))
  except Exception:
    logger.exception('drop_session failed')


reload_admins()
load_settings()
load_sessions()


def audit(actor, action, target_id=None, details=''):
  """سجل العمليات: مين عمل شو ومتى."""
  try:
    actor_id = actor.id if actor else None
    actor_name = display_name(actor) if actor else 'النظام'
    q_exec(
        'INSERT INTO audit_log (ts, actor_id, actor_name, action, target_id,'
        ' details) VALUES (?, ?, ?, ?, ?, ?)',
        (now_utc_str(), actor_id, actor_name, action, target_id, details),
    )
  except Exception:
    logger.exception('audit failed')


# ==================== أدوات الإرسال (تأثيرات + ملصقات) ====================
def send_with_effect(chat_id, text, effect_id=None, **kwargs):
  if effect_id:
    try:
      return send_msg(
          chat_id, text, message_effect_id=effect_id, **kwargs
      )
    except Exception as e:
      logger.info('effect failed (%s), retrying without effect', e)
  return send_msg(chat_id, text, **kwargs)


def send_sticker_slot(chat_id, slot):
  file_id = get_setting(f'sticker_{slot}')
  if not file_id:
    return
  try:
    bot.send_sticker(chat_id, file_id)
  except Exception as e:
    logger.info('sticker %s failed: %s', slot, e)


def show(call, text, markup=None):
  """يعدّل رسالة الأدمن الحالية، وإذا ما أمكن (مثلاً صورة) بيبعت رسالة جديدة."""
  t, ents = rich(text)
  try:
    _premium_guard(
        lambda plain: bot.edit_message_text(
            text if plain else t, call.message.chat.id,
            call.message.message_id,
            **({} if plain or not ents else {'entities': ents}),
            reply_markup=markup),
        ents, markup,
    )
  except Exception as e:
    if 'not modified' in str(e).lower():
      return
    send_msg(call.message.chat.id, text, reply_markup=markup)


def edit_caption_or_text(chat_id, message_id, text, markup=None):
  t, ents = rich(text)
  try:
    _premium_guard(
        lambda plain: bot.edit_message_caption(
            caption=text if plain else t, chat_id=chat_id,
            message_id=message_id,
            **({} if plain or not ents else {'caption_entities': ents}),
            reply_markup=markup),
        ents, markup,
    )
  except Exception:
    try:
      _premium_guard(
          lambda plain: bot.edit_message_text(
              text if plain else t, chat_id, message_id,
              **({} if plain or not ents else {'entities': ents}),
              reply_markup=markup),
          ents, markup,
      )
    except Exception:
      pass


def answer(call, text=None, alert=False):
  try:
    bot.answer_callback_query(call.id, text, show_alert=alert)
  except Exception:
    pass


# ==================== توليد الروابط ====================
def revoke_old_links(target_id):
  rows = q_all(
      'SELECT chat_id, link FROM invite_links WHERE telegram_id = ? AND'
      ' revoked = 0',
      (target_id,),
  )
  for chat_id, link in rows:
    try:
      bot.revoke_chat_invite_link(chat_id, link)
    except Exception as e:
      logger.info('revoke failed %s: %s', link, e)
  if rows:
    q_exec(
        'UPDATE invite_links SET revoked = 1 WHERE telegram_id = ?',
        (target_id,),
    )


def generate_links(target_id):
  """يولّد رابط لكل قناة/مجموعة مضافة. يرجع (links, errors)."""
  channels = q_all('SELECT chat_id, title, kind FROM channels ORDER BY id')
  if not channels:
    return [], ['لا توجد قنوات/مجموعات مضافة. أضفها من ⚙️ الإعدادات.']
  revoke_old_links(target_id)
  expire = int(time.time()) + INVITE_TTL_HOURS * 3600
  links, errors = [], []
  for chat_id, title, kind in channels:
    try:
      link = bot.create_chat_invite_link(
          chat_id=chat_id, member_limit=1, expire_date=expire
      ).invite_link
      links.append((title, kind, link))
      q_exec(
          'INSERT INTO invite_links (telegram_id, chat_id, link, created_at)'
          ' VALUES (?, ?, ?, ?)',
          (target_id, chat_id, link, now_utc_str()),
      )
    except Exception as e:
      errors.append(f'{title}: {e}')
  return links, errors


def links_markup(links):
  markup = types.InlineKeyboardMarkup(row_width=1)
  for title, kind, link in links:
    icon = '📚' if kind == 'channel' else '💬'
    markup.add(btn(f'{icon} {title}', url=link, style='primary'))
  return markup


def send_links_to_student(target_id, full_name, links, first_time=True):
  """يبعت الروابط للطالب. يرمي استثناء لو ما قدر (مثلاً حاجب البوت)."""
  if first_time:
    send_sticker_slot(target_id, 'approved')
    text = (
        f'🎉 مبارك قبولك وتوثيق حسابك يا {full_name}!\n\n'
        f'أهلاً بك رسمياً في قسم {branch_name()} 🎓\n'
        f'ممثل القسم: {owner_tag()}\n\n'
        f'👇 إليك روابط الانضمام الخاصة بدفعتك (سنة أولى مستجدين):\n'
        f'⏰ الروابط صالحة {INVITE_TTL_HOURS} ساعة ولمرة واحدة فقط، ولا تشاركها'
        ' مع أحد.\n🔗 بتقدر تطلب روابط جديدة مرة وحدة بس إذا انتهت صلاحيتها.'
    )
    send_with_effect(
        target_id, text, EFFECT_CONFETTI, reply_markup=links_markup(links)
    )
  else:
    text = (
        f'🔗 روابط الانضمام الجديدة يا {full_name}:\n'
        f'⏰ صالحة {INVITE_TTL_HOURS} ساعة ولمرة واحدة، والروابط القديمة'
        ' ألغيناها.'
    )
    send_msg(target_id, text, reply_markup=links_markup(links))


# ==================== Excel ====================
def build_excel(only_approved=True):
  where = "WHERE status = 'approved'" if only_approved else ''
  rows = q_all(
      'SELECT full_name, phone, username, telegram_id, status, created_at,'
      ' approved_at, approved_by, attempts FROM freshmen'
      f' {where} ORDER BY COALESCE(approved_at, created_at), full_name'
  )
  wb = Workbook()
  ws = wb.active
  ws.title = 'المستجدون المقبولون' if only_approved else 'كل الطلبات'
  ws.sheet_view.rightToLeft = True

  headers = ['م', 'الاسم الكامل', 'رقم الهاتف', 'يوزر تيليجرام', 'Telegram ID']
  if not only_approved:
    headers.append('الحالة')
  headers += ['تاريخ التقديم', 'تاريخ القبول', 'قُبل بواسطة', 'عدد المحاولات']
  ws.append(headers)

  head_fill = PatternFill('solid', fgColor='1F4E78')
  thin = Side(style='thin', color='BFBFBF')
  border = Border(left=thin, right=thin, top=thin, bottom=thin)
  for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = head_fill
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = border

  for i, (name, phone, username, tid, status, created, appr_at, appr_by,
          attempts) in enumerate(rows, 1):
    row = [
        i,
        name or '',
        str(phone or ''),
        f'@{username}' if username else '',
        int(tid),
    ]
    if not only_approved:
      row.append(STATUS_LABELS.get(status, status))
    row += [
        to_local(created),
        to_local(appr_at),
        appr_by or '',
        attempts or 0,
    ]
    ws.append(row)

  phone_col = 3
  id_col = 5
  for r in range(2, ws.max_row + 1):
    # حماية من حقن المعادلات: كل النصوص تتخزن كنص صرف حتى لو بدأت بـ = أو @
    for c in ws[r]:
      if isinstance(c.value, str):
        c.data_type = 's'
    ws.cell(r, phone_col).number_format = '@'
    ws.cell(r, id_col).number_format = '0'
    for c in ws[r]:
      c.border = border
      c.alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(r, 2).alignment = Alignment(horizontal='right', vertical='center')

  for idx in range(1, len(headers) + 1):
    letter = get_column_letter(idx)
    longest = max(
        (len(str(c.value)) for c in ws[letter] if c.value is not None),
        default=8,
    )
    ws.column_dimensions[letter].width = min(max(longest + 4, 10), 40)
  ws.freeze_panes = 'A2'
  ws.auto_filter.ref = ws.dimensions

  out = io.BytesIO()
  wb.save(out)
  return out.getvalue(), len(rows)


def send_excel_to(chat_ids, only_approved=True, caption=None):
  data, count = build_excel(only_approved)
  stamp = datetime.now(LOCAL_TZ).strftime('%Y-%m-%d')
  base = 'Freshmen_Accepted' if only_approved else 'Freshmen_All'
  file_name = f'{base}_{stamp}.xlsx'
  caption = caption or (
      f'📊 قائمة المستجدين المقبولين — العدد: {count}'
      if only_approved
      else f'📋 كل الطلبات وحالاتها — العدد: {count}'
  )
  for cid in chat_ids:
    try:
      cap_t, cap_e = rich(caption)
      bot.send_document(
          cid,
          document=types.InputFile(io.BytesIO(data), file_name=file_name),
          caption=cap_t,
          **({'caption_entities': cap_e} if cap_e else {}),
      )
    except Exception as e:
      logger.warning('فشل إرسال Excel إلى %s: %s', cid, e)
  return count


def auto_excel_after_approval(full_name):
  if get_setting('auto_excel', 'on') != 'on':
    return
  try:
    count = q_one("SELECT COUNT(*) FROM freshmen WHERE status = 'approved'")[0]
    send_excel_to(
        list(_admin_ids),
        True,
        caption=(
            f'✅ تم قبول: {full_name}\n📊 قائمة المستجدين المقبولين'
            f' المحدّثة — العدد: {count}'
        ),
    )
  except Exception:
    logger.exception('auto excel failed')


def export_photos_zip(chat_id):
  """يجمع صور الهوية والمفاضلة للمقبولين بملف/ملفات ZIP (للتقديم للجامعة)."""
  try:
    rows = q_all(
        'SELECT full_name, phone, admission_photo_id, id_photo_id FROM'
        " freshmen WHERE status = 'approved' ORDER BY"
        ' COALESCE(approved_at, created_at), full_name'
    )
    if not rows:
      send_msg(chat_id, '📭 لا يوجد طلاب مقبولون بعد.')
      return
    progress = send_msg(
        chat_id, f'⏳ جاري تجهيز صور {len(rows)} طالب... ممكن ياخد دقائق.'
    )
    state = {'part': 1, 'size': 0, 'buf': io.BytesIO()}
    state['zf'] = zipfile.ZipFile(state['buf'], 'w', zipfile.ZIP_STORED)
    failed = []

    def flush(last=False):
      state['zf'].close()
      if state['size'] > 0:
        bot.send_document(
            chat_id,
            document=types.InputFile(
                io.BytesIO(state['buf'].getvalue()),
                file_name=f"Freshmen_Photos_part{state['part']}.zip",
            ),
            caption=f"🗂 صور الهويات والمفاضلة — جزء {state['part']}",
        )
      if not last:
        state['part'] += 1
        state['size'] = 0
        state['buf'] = io.BytesIO()
        state['zf'] = zipfile.ZipFile(state['buf'], 'w', zipfile.ZIP_STORED)

    for i, (name, phone, adm, idp) in enumerate(rows, 1):
      folder = f'{i:03d}_{safe_filename(name)}'
      for label, file_id in (('ID_الهوية', idp), ('المفاضلة', adm)):
        if not file_id:
          continue
        try:
          info = bot.get_file(file_id)
          data = bot.download_file(info.file_path)
        except Exception as e:
          failed.append(f'{name} ({label})')
          logger.warning('zip download failed %s: %s', name, e)
          continue
        if state['size'] + len(data) > ZIP_PART_LIMIT:
          flush()
        state['zf'].writestr(f'{folder}/{label}.jpg', data)
        state['size'] += len(data)
      if i % 20 == 0:
        try:
          bot.edit_message_text(
              f'⏳ تم تجهيز {i} / {len(rows)}...', chat_id,
              progress.message_id,
          )
        except Exception:
          pass
    flush(last=True)
    msg = '✅ انتهى التجهيز.'
    if failed:
      msg += f'\n⚠️ تعذر تحميل {len(failed)} صورة:\n' + '\n'.join(failed[:15])
    send_msg(chat_id, msg)
  except Exception as e:
    logger.exception('zip export failed')
    send_msg(chat_id, f'❌ فشل تجهيز الملفات: {e}')


def backup_loop():
  """نسخة احتياطية أسبوعية (Excel كامل) للأدمنية، لأن Render ما بيحفظ ملفات."""
  while True:
    try:
      if get_setting('weekly_backup', 'on') == 'on':
        last = get_setting('last_backup', '')
        now = int(time.time())
        if not last:
          set_setting('last_backup', str(now))
        elif now - int(last) >= 7 * 86400:
          total = q_one('SELECT COUNT(*) FROM freshmen')[0]
          if total:
            send_excel_to(
                list(_admin_ids),
                False,
                caption='🗓 نسخة احتياطية أسبوعية لكل الطلبات.',
            )
          set_setting('last_backup', str(now))
    except Exception:
      logger.exception('backup loop error')
    time.sleep(3600)


# ==================== لوحة تحكم الأدمن ====================
def get_admin_keyboard():
  m = types.InlineKeyboardMarkup(row_width=2)
  m.add(
      btn('📊 الإحصائيات', 'admin_stats', style='primary'),
      btn('📁 الأرشيف', 'admin_archive', style='primary'),
  )
  m.add(btn('🔍 بحث عن طالب', 'admin_search'),
        btn('🔓 فك قفل طالب', 'admin_reset'))
  m.add(
      btn('📥 Excel المقبولين', 'admin_export', style='success'),
      btn('📋 Excel الكل', 'admin_export_all'),
  )
  m.add(
      btn('🗂 صور المقبولين ZIP', 'admin_zip', style='success'),
      btn('📣 رسالة جماعية', 'admin_broadcast', style='primary'),
  )
  m.add(btn('📜 سجل العمليات', 'admin_log'),
        btn('⚙️ الإعدادات', 'admin_settings'))
  return m


def admin_panel_text():
  return (
      '🛠️ لوحة تحكم الهيئة المباشرة (Admin Panel)\n'
      f'🎓 {branch_name()}\n\n'
      'جميع الطلبات تصلك هنا في محادثتك المباشرة.\n'
      'اختر من الأزرار أدناه 👇'
  )


@bot.message_handler(commands=['admin'])
def admin_command(message):
  if not is_admin(message.from_user.id):
    reply_msg(message, '⚠️ عذراً، هذه اللوحة مخصصة للأدمن فقط.')
    return
  admin_input_states.pop(message.from_user.id, None)
  send_msg(
      message.chat.id, admin_panel_text(), reply_markup=get_admin_keyboard()
  )


@bot.message_handler(commands=['cancel'])
def cancel_command(message):
  uid = message.from_user.id
  if is_admin(uid):
    if admin_input_states.pop(uid, None):
      reply_msg(message, '↩️ تم إلغاء العملية.')
    return
  if uid in sessions:
    drop_session(uid)
    reply_msg(message, '↩️ تم إلغاء التسجيل. اكتب /start للبدء من جديد.')


def cb_stats(call):
  d = dict(q_all('SELECT status, COUNT(*) FROM freshmen GROUP BY status'))
  total = sum(d.values())
  since = (datetime.now(timezone.utc) - timedelta(days=1)).strftime(
      '%Y-%m-%d %H:%M:%S'
  )
  last_day = q_one(
      'SELECT COUNT(*) FROM freshmen WHERE COALESCE(updated_at, created_at)'
      ' >= ?',
      (since,),
  )[0]
  text = (
      '📊 إحصائيات توثيق المستجدين:\n\n'
      f'👥 إجمالي المتقدمين: {total}\n'
      f'✅ المقبولة: {d.get("approved", 0)}\n'
      f'⏳ المعلقة: {d.get("pending", 0)}\n'
      f'❌ المرفوضة: {d.get("rejected", 0)}\n'
      f'🚫 المحظورة: {d.get("banned", 0)}\n'
      f'🔓 المفتوحة من جديد: {d.get("reset", 0)}\n\n'
      f'🕒 نشاط آخر 24 ساعة: {last_day}'
  )
  show(call, text, get_admin_keyboard())
  answer(call)


def cb_archive(call, status='all', page=0):
  where = '' if status == 'all' else 'WHERE status = ?'
  params = () if status == 'all' else (status,)
  total = q_one(f'SELECT COUNT(*) FROM freshmen {where}', params)[0]
  pages = max((total + PAGE_SIZE - 1) // PAGE_SIZE, 1)
  page = min(max(page, 0), pages - 1)
  rows = q_all(
      'SELECT telegram_id, full_name, status, COALESCE(updated_at,'
      f' created_at) FROM freshmen {where} ORDER BY COALESCE(updated_at,'
      ' created_at) DESC LIMIT ? OFFSET ?',
      params + (PAGE_SIZE, page * PAGE_SIZE),
  )
  markup = types.InlineKeyboardMarkup(row_width=4)
  filters = [('all', 'الكل'), ('approved', '✅'), ('pending', '⏳'),
             ('rejected', '❌')]
  markup.add(*[
      btn(label, f'admin_arch_{key}_0',
          style='primary' if key == status else None)
      for key, label in filters
  ])
  for tid, name, st, ts in rows:
    markup.add(btn(
        f'{status_icon(st)} {(name or "؟")[:28]} ({to_local(ts)[:10]})',
        f'view_file_{tid}',
    ))
  nav = []
  if page > 0:
    nav.append(btn('◀️ السابق', f'admin_arch_{status}_{page - 1}'))
  if page < pages - 1:
    nav.append(btn('التالي ▶️', f'admin_arch_{status}_{page + 1}'))
  if nav:
    markup.add(*nav)
  markup.add(btn('🔙 اللوحة', 'admin_panel'))
  text = (
      f'📁 الأرشيف — العدد: {total} — صفحة {page + 1}/{pages}\n'
      'اضغط على اسم الطالب لفتح مجلده وصوره:'
      if rows
      else '📁 لا يوجد طلاب بهذا الفلتر.'
  )
  show(call, text, markup)
  answer(call)


def cb_log(call):
  rows = q_all(
      'SELECT ts, actor_name, action, target_id, details FROM audit_log'
      ' ORDER BY id DESC LIMIT 15'
  )
  if not rows:
    text = '📜 السجل فارغ.'
  else:
    lines = ['📜 آخر العمليات:\n']
    for ts, actor, action, target, details in rows:
      line = f'• {to_local(ts)} | {actor} | {action}'
      if target:
        line += f' | {target}'
      if details:
        line += f' | {details}'
      lines.append(line)
    text = '\n'.join(lines)[:4000]
  show(call, text, types.InlineKeyboardMarkup().add(
      btn('🔙 اللوحة', 'admin_panel')))
  answer(call)


# ==================== عرض مجلد الطالب للأدمن ====================
def student_actions_markup(tid, status):
  m = types.InlineKeyboardMarkup(row_width=2)
  if status == 'pending':
    m.add(
        btn('✅ قبول وتوليد الرابط', f'approve_{tid}', style='success'),
        btn('❌ رفض الطلب', f'reject_menu_{tid}', style='danger'),
    )
  elif status == 'rejected':
    m.add(btn('✅ قبول (تراجع عن الرفض)', f'approve_{tid}', style='success'))
  elif status == 'approved':
    m.add(btn('🔁 إعادة إرسال الروابط', f'resend_{tid}', style='primary'))
  if status == 'banned':
    m.add(btn('♻️ فك الحظر', f'unban_{tid}', style='success'))
  else:
    m.add(btn('🚫 حظر الطالب', f'ban_{tid}', style='danger'))
  return m


def send_student_file(chat_id, tid):
  row = q_one(
      'SELECT full_name, phone, username, admission_photo_id, id_photo_id,'
      ' status, created_at, attempts, approved_by, approved_at, reject_reason'
      ' FROM freshmen WHERE telegram_id = ?',
      (tid,),
  )
  if not row:
    send_msg(chat_id, '❌ سجل الطالب غير موجود.')
    return
  (full_name, phone, username, adm, idp, status, created, attempts,
   appr_by, appr_at, reason) = row
  info = (
      '📁 مجلد الطالب الرقمي:\n\n'
      f'👤 الاسم: {full_name}\n'
      f'📱 الهاتف: {phone}\n'
      f'🆔 Telegram ID: {tid}\n'
      f'🔗 اليوزر: {"@" + username if username else "لا يوجد"}\n'
      f'📌 الحالة: {status_label(status)}\n'
      f'🔁 عدد المحاولات: {attempts or 0}\n'
      f'🕒 تاريخ التقديم: {to_local(created)}'
  )
  if appr_by:
    info += f'\n✅ قُبل بواسطة: {appr_by} ({to_local(appr_at)})'
  if reason and status == 'rejected':
    info += f'\n📝 سبب الرفض: {reason}'
  send_msg(chat_id, info)
  if adm:
    try:
      send_photo_rich(chat_id, adm, caption=f'📄 صورة المفاضلة: {full_name}')
    except Exception:
      pass
  markup = student_actions_markup(tid, status)
  if idp:
    try:
      send_photo_rich(
          chat_id, idp, caption=f'🪪 صورة الهوية: {full_name}',
          reply_markup=markup,
      )
      return
    except Exception:
      pass
  send_msg(chat_id, '⚙️ إجراءات:', reply_markup=markup)


# ==================== القبول والرفض (ذرّي) ====================
def claimed(rowcount, tid, expect_status, expect_field=None, expect_value=None):
  """يتأكد إنو هالأدمن هو اللي نفّذ التحديث فعلاً (حماية من ضغطتين متزامنتين)."""
  if rowcount is not None and rowcount >= 0:
    return rowcount > 0
  row = q_one(
      f'SELECT status, {expect_field or "status"} FROM freshmen WHERE'
      ' telegram_id = ?',
      (tid,),
  )
  return bool(row) and row[0] == expect_status and (
      expect_field is None or row[1] == expect_value
  )


def who_handled(tid):
  row = q_one(
      'SELECT status, approved_by FROM freshmen WHERE telegram_id = ?', (tid,)
  )
  if not row:
    return 'غير موجود'
  st, by = row
  return f'{status_label(st)}' + (f' بواسطة {by}' if by else '')


def do_approve(target_id, admin_user, chat_id=None, message_id=None):
  """يرجع (ok, text)."""
  row = q_one(
      'SELECT full_name, phone, status FROM freshmen WHERE telegram_id = ?',
      (target_id,),
  )
  if not row:
    return False, '❌ لم يتم العثور على الطالب.'
  full_name, phone, old_status = row
  if old_status == 'approved':
    return False, f'ℹ️ هذا الطالب مقبول مسبقاً ({who_handled(target_id)}).'
  if old_status not in ('pending', 'rejected'):
    return False, f'⚠️ لا يمكن قبول طالب حالته: {status_label(old_status)}'

  admin_name = display_name(admin_user)
  now = now_utc_str()
  rc = q_exec(
      "UPDATE freshmen SET status = 'approved', approved_by = ?, approved_at"
      ' = ?, reject_reason = NULL, updated_at = ? WHERE telegram_id = ? AND'
      ' status = ?',
      (admin_name, now, now, target_id, old_status),
  )
  if not claimed(rc, target_id, 'approved', 'approved_by', admin_name):
    return False, f'ℹ️ تم التعامل مع الطلب مسبقاً: {who_handled(target_id)}'

  links, errors = generate_links(target_id)
  if not links:
    q_exec(
        'UPDATE freshmen SET status = ?, approved_by = NULL, approved_at ='
        " NULL WHERE telegram_id = ? AND status = 'approved'",
        (old_status, target_id),
    )
    return False, (
        '❌ تعذر توليد الروابط:\n' + '\n'.join(errors) +
        '\n\nتأكد أن البوت مشرف في القنوات/المجموعات ولديه صلاحية دعوة'
        ' المستخدمين (Invite Users).'
    )

  q_exec(
      "UPDATE submissions SET result = 'approved' WHERE id = (SELECT MAX(id)"
      ' FROM submissions WHERE telegram_id = ?)',
      (target_id,),
  )
  audit(admin_user, 'قبول', target_id, full_name)

  text = f'✅ تم قبول الطالب ({full_name} - {phone}) وإرسال الروابط له!'
  try:
    send_links_to_student(target_id, full_name, links, first_time=True)
  except Exception as e:
    text = (
        f'⚠️ تم قبول ({full_name} - {phone}) لكن تعذر إرسال الروابط له'
        f' (غالباً حاجب البوت): {e}\nاستخدم زر 🔁 إعادة إرسال الروابط لاحقاً.'
    )
  if errors:
    text += '\n⚠️ تعذر توليد روابط لـ:\n' + '\n'.join(errors)

  Thread(
      target=auto_excel_after_approval, args=(full_name,), daemon=True
  ).start()
  return True, text


def do_reject(target_id, admin_user, reason_text):
  row = q_one(
      'SELECT full_name, status, attempts FROM freshmen WHERE telegram_id = ?',
      (target_id,),
  )
  if not row:
    return False, '❌ لم يتم العثور على الطالب.'
  full_name, old_status, attempts = row
  if old_status != 'pending':
    return False, f'ℹ️ تم التعامل مع الطلب مسبقاً: {who_handled(target_id)}'
  rc = q_exec(
      "UPDATE freshmen SET status = 'rejected', reject_reason = ?, updated_at"
      " = ? WHERE telegram_id = ? AND status = 'pending'",
      (reason_text, now_utc_str(), target_id),
  )
  if not claimed(rc, target_id, 'rejected', 'reject_reason', reason_text):
    return False, f'ℹ️ تم التعامل مع الطلب مسبقاً: {who_handled(target_id)}'
  q_exec(
      "UPDATE submissions SET result = 'rejected', reason = ? WHERE id ="
      ' (SELECT MAX(id) FROM submissions WHERE telegram_id = ?)',
      (reason_text, target_id),
  )
  audit(admin_user, 'رفض', target_id, f'{full_name}: {reason_text}')

  left = max(MAX_ATTEMPTS - (attempts or 0), 0)
  markup = types.InlineKeyboardMarkup()
  if left > 0:
    markup.add(btn('🔄 إعادة رفع الصور والبيانات', 'retry_freshman',
                   style='primary'))
  note = (
      f'\n\n🔁 باقي لك {left} محاولة.' if left > 0
      else f'\n\n⛔ استنفدت المحاولات، تواصل مع {owner_tag()}.'
  )
  text = f'✅ تم رفض الطلب.\nالسبب: {reason_text}'
  try:
    send_sticker_slot(target_id, 'rejected')
    send_msg(
        target_id,
        f'❌ عذراً، تعذر قبول طلب التوثيق الخاص بك.\n\n📌 السبب:'
        f' {reason_text}{note}',
        reply_markup=markup if left > 0 else None,
    )
  except Exception as e:
    text += f'\n⚠️ تعذر إبلاغ الطالب (غالباً حاجب البوت): {e}'
  return True, text


REJECT_REASONS = {
    '1': '📷 الصور المرفقة غير واضحة، يرجى إعادة التصوير بشكل جلي والإرسال مجدداً.',
    '2': '📄 بطاقة المفاضلة المرفقة لا توضح القبول في قسم الهندسة الإلكترونية والاتصالات.',
    '3': '🪪 صورة الهوية الشخصية غير واضحة أو غير مطابقة للبيانات.',
}


def reject_menu_markup(target_id):
  m = types.InlineKeyboardMarkup(row_width=1)
  m.add(
      btn('📷 الصورة غير واضحة', f'reject_reason_{target_id}_1'),
      btn('📄 المفاضلة لا تحتوي القسم', f'reject_reason_{target_id}_2'),
      btn('🪪 صورة الهوية غير مطابقة', f'reject_reason_{target_id}_3'),
      btn('✍️ سبب آخر (أكتبه)', f'reject_reason_{target_id}_c',
          style='primary'),
      btn('↩️ تراجع', f'view_file_{target_id}'),
  )
  return m


# ==================== بدء التسجيل للمستجدين ====================
def progress_bar(step):
  return '🟩' * step + '⬜' * (4 - step)


def begin_registration_flow(user_id, chat_id):
  row = q_one(
      'SELECT status, full_name, attempts, links_count FROM freshmen WHERE'
      ' telegram_id = ?',
      (user_id,),
  )
  if row:
    status, full_name, attempts, links_used = row
    if status == 'banned':
      send_msg(
          chat_id, f'🚫 تم إيقاف حسابك عن التوثيق. تواصل مع {owner_tag()}.'
      )
      return
    if status == 'approved':
      can_regen = MAX_LINK_REGEN > 0 and (links_used or 0) < MAX_LINK_REGEN
      markup = types.InlineKeyboardMarkup()
      if can_regen:
        markup.add(btn('🔗 احصل على روابط جديدة (مرة واحدة)', 'mylinks',
                       style='primary'))
      send_msg(
          chat_id,
          f'🎉 أهلاً بك مجدداً يا {full_name}!\n\nتم توثيق حسابك مسبقاً.'
          + ('\nإذا انتهت صلاحية روابطك اضغط الزر:' if can_regen else
             f'\nإذا تحتاج روابط جديدة تواصل مع {owner_tag()}.'),
          reply_markup=markup if can_regen else None,
      )
      return
    if status == 'pending':
      send_msg(
          chat_id,
          '⏳ طلبك قيد المراجعة حالياً من قبل الهيئة.\n\nيرجى الانتظار، وسيصلك'
          ' رابط الانضمام هنا فور التدقيق والقبول.',
      )
      return
    if status == 'rejected' and (attempts or 0) >= MAX_ATTEMPTS:
      send_msg(
          chat_id,
          f'⛔ استنفدت عدد محاولات التقديم. تواصل مع {owner_tag()}.',
      )
      return

  save_session(user_id, step='NAME', name=None, phone=None,
               admission_photo=None)
  send_sticker_slot(chat_id, 'welcome')
  welcome_text = (
      f'🎓 أهلاً بك في بوت توثيق الطلاب المستجدين\n'
      f'قسم {branch_name()} - الهمك\n'
      f'( بَرمَجَ هذا البوت {owner_tag()} )\n\n'
      'للانضمام لقناة ومجموعة السنة الأولى، يرجى إكمال خطوات التوثيق للتأكد'
      ' من قبولك بالقسم.\n\n'
      f'{progress_bar(0)}\n'
      '✍️ الخطوة (1/4): يرجى كتابة اسمك الثلاثي الكامل:'
  )
  send_msg(chat_id, welcome_text,
                   reply_markup=types.ReplyKeyboardRemove())


@bot.message_handler(commands=['start'])
def start_cmd(message):
  if message.chat.type != 'private':
    return
  user_id = message.from_user.id
  if is_admin(user_id):
    admin_command(message)
    return
  begin_registration_flow(user_id, message.chat.id)


# ==================== استقبال مدخلات الأدمن (بحث، إعدادات، ...) ====================
ALL_CONTENT = [
    'text', 'photo', 'video', 'document', 'audio', 'voice', 'sticker',
    'animation', 'video_note', 'contact', 'location',
]


def admin_state_filter(msg):
  return (
      msg.chat.type == 'private'
      and is_admin(msg.from_user.id)
      and msg.from_user.id in admin_input_states
      and not (msg.text or '').startswith('/')
  )


def set_state(admin_id, state, **data):
  admin_input_states[admin_id] = {'state': state, 'data': data}


def extract_forward_chat(m):
  fc = getattr(m, 'forward_from_chat', None)
  if fc:
    return fc
  fo = getattr(m, 'forward_origin', None)
  return getattr(fo, 'chat', None) or getattr(fo, 'sender_chat', None)


def extract_forward_user(m):
  fu = getattr(m, 'forward_from', None)
  if fu:
    return fu
  fo = getattr(m, 'forward_origin', None)
  return getattr(fo, 'sender_user', None)


@bot.message_handler(func=admin_state_filter, content_types=ALL_CONTENT)
def handle_admin_input(message):
  admin_id = message.from_user.id
  st = admin_input_states.get(admin_id)
  if not st:
    return
  state, data = st['state'], st['data']
  text = (message.text or '').strip()

  if state == 'awaiting_search':
    if not text:
      reply_msg(message, '✍️ أرسل نص للبحث (اسم / هاتف / ID).')
      return
    like = f'%{text}%'
    phone_like = f'%{clean_phone(text)}%' if clean_phone(text) else like
    rows = q_all(
        'SELECT telegram_id, full_name, phone, status FROM freshmen WHERE'
        ' phone LIKE ? OR phone LIKE ? OR CAST(telegram_id AS TEXT) LIKE ?'
        ' OR full_name LIKE ? LIMIT 15',
        (like, phone_like, like, like),
    )
    admin_input_states.pop(admin_id, None)
    if not rows:
      reply_msg(message, '❌ لم يتم العثور على أي طالب مطابق للبحث.')
      return
    markup = types.InlineKeyboardMarkup()
    for tid, name, phone, status in rows:
      markup.add(btn(f'{status_icon(status)} {name} ({phone})',
                     f'view_file_{tid}'))
    reply_msg(message, f'🔍 نتائج البحث عن ({text}):', reply_markup=markup)

  elif state == 'awaiting_reset':
    phone = clean_phone(text)
    tid = int(text) if text.isdigit() and len(text) > 8 and not phone.startswith('0') else None
    row = q_one(
        'SELECT telegram_id, full_name, status FROM freshmen WHERE phone = ?'
        ' OR telegram_id = ?',
        (phone, tid if tid is not None else -1),
    )
    admin_input_states.pop(admin_id, None)
    if not row:
      reply_msg(message, '❌ لم يتم العثور على الطالب.')
      return
    target, name, old = row
    if old == 'banned':
      reply_msg(message, '🚫 هذا الطالب محظور، افتح مجلده وفك الحظر.')
      return
    q_exec(
        "UPDATE freshmen SET status = 'reset', attempts = 0, approved_by ="
        ' NULL, approved_at = NULL, updated_at = ? WHERE telegram_id = ?',
        (now_utc_str(), target),
    )
    revoke_old_links(target)
    drop_session(target)
    audit(message.from_user, 'فك قفل', target, name)
    reply_msg(message, f'✅ تم فك القفل عن ({name}) وصار بإمكانه التسجيل من جديد.')

  elif state == 'awaiting_reject_text':
    if not text:
      reply_msg(message, '✍️ اكتب سبب الرفض كنص.')
      return
    admin_input_states.pop(admin_id, None)
    ok, msg = do_reject(data['target'], message.from_user, text[:300])
    reply_msg(message, msg)
    if ok and data.get('chat_id'):
      edit_caption_or_text(data['chat_id'], data['message_id'], msg)

  elif state == 'awaiting_broadcast':
    admin_input_states.pop(admin_id, None)
    count = q_one("SELECT COUNT(*) FROM freshmen WHERE status = 'approved'")[0]
    try:
      bot.copy_message(message.chat.id, message.chat.id, message.message_id)
    except Exception:
      pass
    set_state(admin_id, 'broadcast_ready', chat_id=message.chat.id,
              message_id=message.message_id)
    m = types.InlineKeyboardMarkup(row_width=2)
    m.add(btn(f'✅ إرسال لـ {count} طالب', 'bc_yes', style='success'),
          btn('❌ إلغاء', 'bc_no', style='danger'))
    send_msg(
        message.chat.id,
        f'👆 هذه معاينة الرسالة. بدك ترسلها لكل المقبولين ({count})؟',
        reply_markup=m,
    )

  elif state == 'awaiting_branch':
    if not text:
      return
    set_setting('branch_name', text[:80])
    audit(message.from_user, 'تغيير اسم الفرع', None, text[:80])
    admin_input_states.pop(admin_id, None)
    reply_msg(message, f'✅ تم ضبط اسم الفرع: {text[:80]}')

  elif state == 'awaiting_owner':
    username = text.lstrip('@').strip()
    if not re.fullmatch(r'[A-Za-z0-9_]{4,32}', username):
      reply_msg(message, '⚠️ يوزر غير صالح. أرسله بدون مسافات، مثال: Youssef_Sabra')
      return
    set_setting('owner_username', username)
    audit(message.from_user, 'تغيير يوزر المسؤول', None, username)
    admin_input_states.pop(admin_id, None)
    reply_msg(message, f'✅ تم ضبط يوزر المسؤول: @{username}')

  elif state == 'awaiting_ch_add':
    fc = extract_forward_chat(message)
    if fc:
      chat_id = fc.id
    elif re.fullmatch(r'-?\d{6,}', text):
      chat_id = int(text)
    else:
      reply_msg(message,
          '⚠️ حوّل لي رسالة من القناة/المجموعة، أو أرسل الـ ID الخاص فيها'
          ' (مثال: -1001234567890).',
      )
      return
    try:
      chat = bot.get_chat(chat_id)
      me = bot.get_me()
      member = bot.get_chat_member(chat_id, me.id)
      if member.status not in ('administrator', 'creator'):
        raise RuntimeError('البوت مو مشرف بهالقناة/المجموعة')
      if member.status == 'administrator' and getattr(
          member, 'can_invite_users', True) is False:
        raise RuntimeError('البوت ما عنده صلاحية Invite Users')
    except Exception as e:
      reply_msg(message,
          f'❌ تعذر الإضافة: {e}\nاجعل البوت مشرفاً مع صلاحية دعوة المستخدمين'
          ' ثم جرّب مرة ثانية.',
      )
      return
    kind = 'channel' if chat.type == 'channel' else 'group'
    try:
      q_exec('INSERT INTO channels (chat_id, title, kind) VALUES (?, ?, ?)',
             (chat.id, chat.title or str(chat.id), kind))
    except Exception:
      admin_input_states.pop(admin_id, None)
      reply_msg(message, 'ℹ️ هذه القناة/المجموعة مضافة مسبقاً.')
      return
    audit(message.from_user, 'إضافة قناة', chat.id, chat.title)
    admin_input_states.pop(admin_id, None)
    reply_msg(message, f'✅ تمت إضافة: {chat.title}')

  elif state == 'awaiting_ch_rename':
    if not text:
      return
    q_exec('UPDATE channels SET title = ? WHERE id = ?',
           (text[:60], data['id']))
    audit(message.from_user, 'تعديل اسم قناة', data['id'], text[:60])
    admin_input_states.pop(admin_id, None)
    reply_msg(message, f'✅ تم تغيير الاسم إلى: {text[:60]}')

  elif state == 'awaiting_admin_add':
    new_id = None
    fu = extract_forward_user(message)
    if fu:
      new_id = fu.id
    elif text.isdigit():
      new_id = int(text)
    if not new_id:
      reply_msg(message, '⚠️ أرسل Telegram ID رقمي، أو حوّل رسالة من الشخص.')
      return
    q_exec(
        'INSERT INTO admins (telegram_id, name, added_by, added_at) VALUES'
        ' (?, ?, ?, ?) ON CONFLICT(telegram_id) DO NOTHING',
        (new_id, text if not fu else display_name(fu), admin_id,
         now_utc_str()),
    )
    reload_admins()
    audit(message.from_user, 'إضافة أدمن', new_id)
    admin_input_states.pop(admin_id, None)
    reply_msg(message, f'✅ تمت إضافة الأدمن {new_id}.\n(لازم يعمل /start للبوت مرة وحدة.)')

  elif state == 'awaiting_premium_emoji':
    ents = [e for e in (message.entities or [])
            if getattr(e, 'type', '') == 'custom_emoji']
    if not ents:
      reply_msg(message, '⚠️ ما لقيت إيموجي مميز بالرسالة. اختار من الإيموجيات'
                ' المميزة (اللي بتظهر بحزم Premium) وأرسلها.')
      return
    mp = dict(pe_data()[0])
    raw16 = message.text.encode('utf-16-le')
    added = []
    for e in ents:
      base = raw16[e.offset * 2:(e.offset + e.length) * 2].decode('utf-16-le')
      key = norm_emoji(base)
      if key:
        mp[key] = e.custom_emoji_id
        added.append(key)
    set_setting('pe_map', json.dumps(mp, ensure_ascii=False))
    send_msg(message.chat.id,
             f'✅ تم ربط {len(set(added))}: {" ".join(dict.fromkeys(added))}'
             f'\nالمجموع: {len(mp)}. أرسل المزيد أو اضغط تم.')

  elif state == 'awaiting_sticker':
    stk = message.sticker
    if not stk:
      reply_msg(message, '⚠️ أرسل ملصق (Sticker) تحديداً.')
      return
    set_setting(f'sticker_{data["slot"]}', stk.file_id)
    audit(message.from_user, 'ضبط ملصق', None, data['slot'])
    admin_input_states.pop(admin_id, None)
    reply_msg(message, '✅ تم حفظ الملصق.')


# ==================== خطوات إدخال بيانات الطالب ====================
def in_step(step):
  return lambda msg: (
      msg.chat.type == 'private' and session_step(msg.from_user.id) == step
  )


NAME_RE = re.compile(r'^[\u0600-\u06FFa-zA-Z\s]+$')


@bot.message_handler(content_types=['text'], func=in_step('NAME'))
def handle_name(message):
  user_id = message.from_user.id
  full_name = re.sub(r'\s+', ' ', message.text.strip())
  words = full_name.split(' ')
  if (not NAME_RE.match(full_name) or len(words) < 2 or len(full_name) < 6
      or len(full_name) > 60):
    send_msg(
        message.chat.id,
        '⚠️ اكتب اسمك الثلاثي الكامل بالحروف فقط (بدون أرقام أو رموز).',
    )
    return
  save_session(user_id, step='PHONE', name=full_name)

  markup = types.ReplyKeyboardMarkup(
      row_width=1, resize_keyboard=True, one_time_keyboard=True
  )
  markup.add(kbtn('📱 مشاركة رقم الهاتف لتأكيد الهوية', request_contact=True,
                  style='primary'))
  send_msg(
      message.chat.id,
      f'أهلاً بك يا {full_name}! 👋\n\n{progress_bar(1)}\n'
      '📱 الخطوة (2/4): يرجى الضغط على الزر أدناه لمشاركة رقم هاتفك المعتمد:',
      reply_markup=markup,
  )


@bot.message_handler(content_types=['contact'], func=in_step('PHONE'))
def handle_phone(message):
  if not message.contact:
    return
  user_id = message.from_user.id
  if message.contact.user_id != user_id:
    send_msg(
        message.chat.id,
        '⚠️ تنبيه: يرجى مشاركة رقم الهاتف الخاص بحسابك الحالي حصراً بالضغط على'
        ' الزر.',
    )
    return
  phone = clean_phone(message.contact.phone_number)
  if not (8 <= len(phone) <= 15):
    send_msg(message.chat.id, '⚠️ رقم الهاتف غير صالح، جرّب مجدداً.')
    return
  save_session(user_id, step='ADMISSION_PHOTO', phone=phone)
  send_msg(
      message.chat.id,
      f'✅ تم التحقق من رقم الهاتف بنجاح!\n\n{progress_bar(2)}\n'
      '📄 الخطوة (3/4): يرجى إرسال صورة بطاقة المفاضلة (التي تظهر اسمك وقبولك'
      ' في القسم):',
      reply_markup=types.ReplyKeyboardRemove(),
  )


@bot.message_handler(content_types=['photo'], func=in_step('ADMISSION_PHOTO'))
def handle_admission_photo(message):
  user_id = message.from_user.id
  save_session(
      user_id, step='ID_PHOTO', admission_photo=message.photo[-1].file_id
  )
  send_msg(
      message.chat.id,
      f'✅ تم استلام صورة المفاضلة بنجاح!\n\n{progress_bar(3)}\n'
      '🪪 الخطوة (4/4) والأخيرة: يرجى إرسال صورة البطاقة الشخصية (الهوية):',
  )


@bot.message_handler(content_types=['photo'], func=in_step('ID_PHOTO'))
def handle_id_photo(message):
  user_id = message.from_user.id
  id_photo_id = message.photo[-1].file_id
  s = sessions.get(user_id)
  if not s or not s.get('name') or not s.get('phone') or not s.get(
      'admission_photo'):
    drop_session(user_id)
    send_msg(message.chat.id, '⚠️ صار خلل بالجلسة، اكتب /start من جديد.')
    return
  full_name, phone, admission_photo_id = (
      s['name'], s['phone'], s['admission_photo'])

  cur = q_one('SELECT status FROM freshmen WHERE telegram_id = ?', (user_id,))
  if cur and cur[0] in ('approved', 'pending', 'banned'):
    drop_session(user_id)
    send_msg(
        message.chat.id,
        'ℹ️ حالة حسابك تغيّرت، اكتب /start لمعرفة وضعك الحالي.',
    )
    return

  username = message.from_user.username
  now = now_utc_str()
  q_exec(
      'INSERT INTO freshmen (telegram_id, full_name, phone, username,'
      ' admission_photo_id, id_photo_id, status, attempts, updated_at)'
      " VALUES (?, ?, ?, ?, ?, ?, 'pending', 1, ?) ON CONFLICT(telegram_id)"
      ' DO UPDATE SET full_name = excluded.full_name, phone ='
      ' excluded.phone, username = excluded.username, admission_photo_id ='
      ' excluded.admission_photo_id, id_photo_id = excluded.id_photo_id,'
      " status = 'pending', attempts = COALESCE(freshmen.attempts, 0) + 1,"
      ' updated_at = excluded.updated_at, reject_reason = NULL, approved_by'
      ' = NULL, approved_at = NULL',
      (user_id, full_name, phone, username, admission_photo_id, id_photo_id,
       now),
  )
  q_exec(
      'INSERT INTO submissions (telegram_id, full_name, phone,'
      ' admission_photo_id, id_photo_id, submitted_at) VALUES (?, ?, ?, ?, ?,'
      ' ?)',
      (user_id, full_name, phone, admission_photo_id, id_photo_id, now),
  )
  drop_session(user_id)

  send_sticker_slot(message.chat.id, 'submitted')
  send_msg(
      message.chat.id,
      f'🚀 تم استلام بياناتك وأوراقك بنجاح!\n\n{progress_bar(4)}\n'
      'طلبك الآن قيد المراجعة والتدقيق، وسيصلك إشعار بالقبول مع روابط القنوات'
      ' فور الاعتماد.',
  )

  # كشف التكرار: نفس الرقم أو الاسم عند حساب آخر مقبول/معلق
  dups = q_all(
      'SELECT telegram_id, full_name, status FROM freshmen WHERE telegram_id'
      " != ? AND (phone = ? OR full_name = ?) AND status IN ('approved',"
      " 'pending')",
      (user_id, phone, full_name),
  )
  dup_text = ''
  if dups:
    dup_text = '\n\n⚠️ تنبيه تكرار! نفس الاسم/الرقم موجود عند:\n' + '\n'.join(
        f'• {n} (ID {t}) — {status_label(st)}' for t, n, st in dups
    )
  attempts = q_one(
      'SELECT attempts FROM freshmen WHERE telegram_id = ?', (user_id,)
  )[0]
  username_str = f'@{username}' if username else 'لا يوجد'

  markup = types.InlineKeyboardMarkup(row_width=2)
  markup.add(
      btn('✅ قبول وتوليد الرابط', f'approve_{user_id}', style='success'),
      btn('❌ رفض الطلب', f'reject_menu_{user_id}', style='danger'),
  )

  def send_to_admin(admin_id):
    send_photo_rich(
        admin_id,
        admission_photo_id,
        caption=(
            '📥 طلب توثيق مستجد جديد (1/2 - المفاضلة):\n\n'
            f'👤 الاسم: {full_name}\n'
            f'📱 الهاتف: {phone}\n'
            f'🆔 Telegram ID: {user_id}\n'
            f'👤 المعرف: {username_str}\n'
            f'🔁 المحاولة رقم: {attempts}{dup_text}'
        ),
    )
    send_photo_rich(
        admin_id,
        id_photo_id,
        caption=f'🪪 (2/2 - صورة الهوية) للطالب: {full_name}\n👇 اتخذ القرار بضغطة زر:',
        reply_markup=markup,
    )

  notify_all_admins(send_to_admin)


STEP_HINTS = {
    'NAME': '✍️ اكتب اسمك الثلاثي كنص.',
    'PHONE': '📱 اضغط زر "مشاركة رقم الهاتف" الظاهر تحت.',
    'ADMISSION_PHOTO': '📄 أرسل صورة بطاقة المفاضلة كـ صورة (Photo) وليس ملف.',
    'ID_PHOTO': '🪪 أرسل صورة الهوية كـ صورة (Photo) وليس ملف.',
}


@bot.message_handler(
    content_types=ALL_CONTENT,
    func=lambda m: m.chat.type == 'private'
    and not is_admin(m.from_user.id)
    and session_step(m.from_user.id) is not None,
)
def handle_wrong_input(message):
  step = session_step(message.from_user.id)
  send_msg(
      message.chat.id,
      f'⚠️ هذا غير مطلوب الآن.\n{STEP_HINTS.get(step, "")}\n\nللإلغاء: /cancel',
  )


# ==================== أوامر الأدمن اليدوية ====================
@bot.message_handler(commands=['approve_manual'])
def manual_approve(message):
  if not is_admin(message.from_user.id):
    return
  args = message.text.split()
  if len(args) < 2 or not args[1].strip().isdigit():
    reply_msg(message, '⚠️ اكتب الأمر هكذا:\n/approve_manual TELEGRAM_ID')
    return
  ok, msg = do_approve(int(args[1].strip()), message.from_user)
  reply_msg(message, msg)


# ==================== معالجة الأزرار ====================
def cb_retry(call):
  uid = call.from_user.id
  row = q_one('SELECT status FROM freshmen WHERE telegram_id = ?', (uid,))
  if row and row[0] in ('approved', 'pending', 'banned'):
    answer(call, 'ℹ️ لا يمكنك إعادة التقديم بحالتك الحالية.', True)
    return
  answer(call, '👍 تم فتح التسجيل مجدداً.')
  begin_registration_flow(uid, call.message.chat.id)


def cb_mylinks(call):
  uid = call.from_user.id
  row = q_one(
      'SELECT status, full_name, links_count FROM freshmen WHERE telegram_id'
      ' = ?',
      (uid,),
  )
  if not row or row[0] != 'approved':
    answer(call, '⚠️ حسابك غير موثّق.', True)
    return
  status, full_name, count = row
  if MAX_LINK_REGEN <= 0 or (count or 0) >= MAX_LINK_REGEN:
    answer(call, f'⛔ استخدمت طلب الروابط مسبقاً. تواصل مع {owner_tag()}.',
           True)
    return
  answer(call, '⏳ جاري توليد روابط جديدة...')
  links, errors = generate_links(uid)
  if not links:
    send_msg(call.message.chat.id,
                     f'❌ تعذر توليد الروابط، تواصل مع {owner_tag()}.')
    return
  q_exec('UPDATE freshmen SET links_count = COALESCE(links_count, 0) + 1'
         ' WHERE telegram_id = ?', (uid,))
  send_links_to_student(uid, full_name, links, first_time=False)


def cb_approve(call, target_id):
  ok, msg = do_approve(target_id, call.from_user)
  if ok:
    edit_caption_or_text(call.message.chat.id, call.message.message_id, msg)
    answer(call, '✅ تم')
  else:
    answer(call, msg[:190], True)
    if 'مسبقاً' in msg:
      edit_caption_or_text(call.message.chat.id, call.message.message_id, msg)


def cb_reject_reason(call, target_id, code):
  if code == 'c':
    set_state(call.from_user.id, 'awaiting_reject_text', target=target_id,
              chat_id=call.message.chat.id,
              message_id=call.message.message_id)
    send_msg(call.message.chat.id,
                     '✍️ اكتب سبب الرفض ليتم إبلاغ الطالب فيه (أو /cancel):')
    answer(call)
    return
  reason = REJECT_REASONS.get(code, 'الصور غير مطابقة للشروط.')
  ok, msg = do_reject(target_id, call.from_user, reason)
  if ok:
    edit_caption_or_text(call.message.chat.id, call.message.message_id, msg)
    answer(call, '✅ تم')
  else:
    answer(call, msg[:190], True)
    edit_caption_or_text(call.message.chat.id, call.message.message_id, msg)


def cb_resend(call, target_id):
  row = q_one('SELECT status, full_name FROM freshmen WHERE telegram_id = ?',
              (target_id,))
  if not row or row[0] != 'approved':
    answer(call, '⚠️ الطالب غير مقبول.', True)
    return
  links, errors = generate_links(target_id)
  if not links:
    answer(call, '❌ تعذر توليد الروابط: ' + '; '.join(errors)[:150], True)
    return
  try:
    send_links_to_student(target_id, row[1], links, first_time=False)
    audit(call.from_user, 'إعادة إرسال روابط', target_id, row[1])
    answer(call, '✅ تم إرسال روابط جديدة للطالب.', True)
  except Exception as e:
    answer(call, f'❌ تعذر الإرسال للطالب: {e}'[:190], True)


def cb_ban(call, target_id, ban=True):
  row = q_one('SELECT full_name FROM freshmen WHERE telegram_id = ?',
              (target_id,))
  if not row:
    answer(call, '❌ غير موجود.', True)
    return
  if ban:
    q_exec("UPDATE freshmen SET status = 'banned', updated_at = ? WHERE"
           ' telegram_id = ?', (now_utc_str(), target_id))
    revoke_old_links(target_id)
    drop_session(target_id)
    audit(call.from_user, 'حظر', target_id, row[0])
    answer(call, '🚫 تم حظر الطالب.', True)
  else:
    q_exec("UPDATE freshmen SET status = 'reset', attempts = 0, updated_at ="
           ' ? WHERE telegram_id = ?', (now_utc_str(), target_id))
    audit(call.from_user, 'فك حظر', target_id, row[0])
    answer(call, '♻️ تم فك الحظر، صار بإمكانه التسجيل من جديد.', True)


def cb_broadcast_confirm(call):
  st = admin_input_states.pop(call.from_user.id, None)
  if not st or st['state'] != 'broadcast_ready':
    answer(call, '⚠️ انتهت صلاحية الرسالة، ابدأ من جديد.', True)
    return
  src_chat, src_msg = st['data']['chat_id'], st['data']['message_id']
  admin_chat = call.message.chat.id
  answer(call, '⏳ جاري الإرسال...')
  actor = call.from_user

  def worker():
    ids = [r[0] for r in q_all(
        "SELECT telegram_id FROM freshmen WHERE status = 'approved'")]
    sent = failed = 0
    for tid in ids:
      try:
        bot.copy_message(tid, src_chat, src_msg)
        sent += 1
      except Exception:
        failed += 1
      time.sleep(0.07)
    audit(actor, 'رسالة جماعية', None, f'نجح {sent} / فشل {failed}')
    send_msg(admin_chat,
                     f'📣 انتهى الإرسال.\n✅ وصلت: {sent}\n❌ فشلت: {failed}')

  Thread(target=worker, daemon=True).start()
  edit_caption_or_text(admin_chat, call.message.message_id,
                       '📣 جاري إرسال الرسالة للمقبولين...')


# ---- الإعدادات ----
def settings_markup():
  auto = get_setting('auto_excel', 'on') == 'on'
  weekly = get_setting('weekly_backup', 'on') == 'on'
  m = types.InlineKeyboardMarkup(row_width=1)
  m.add(
      btn('🏷 اسم الفرع', 'set_branch'),
      btn('👤 يوزر المسؤول', 'set_owner'),
      btn('📡 القنوات والمجموعات', 'set_channels', style='primary'),
      btn('👥 الأدمنية', 'set_admins'),
      btn('🌟 الإيموجي المميز (Premium)', 'set_pe', style='primary'),
      btn('🎭 الملصقات', 'set_stickers'),
      btn(f'📎 Excel تلقائي مع كل قبول: {"✅ مفعّل" if auto else "⛔ موقّف"}',
          'set_autoxl', style='success' if auto else 'danger'),
      btn(f'🗓 نسخة احتياطية أسبوعية: {"✅ مفعّلة" if weekly else "⛔ موقّفة"}',
          'set_backup', style='success' if weekly else 'danger'),
      btn('🔙 اللوحة', 'admin_panel'),
  )
  return m


def settings_text():
  return (
      '⚙️ الإعدادات\n\n'
      f'🏷 اسم الفرع: {branch_name()}\n'
      f'👤 المسؤول: {owner_tag()}\n'
      f'⏰ صلاحية الروابط: {INVITE_TTL_HOURS} ساعة'
  )


def channels_view(call):
  rows = q_all('SELECT id, title, kind FROM channels ORDER BY id')
  m = types.InlineKeyboardMarkup(row_width=1)
  for cid, title, kind in rows:
    m.add(btn(f'{"📚" if kind == "channel" else "💬"} {title}', f'set_ch_{cid}'))
  m.add(btn('➕ إضافة قناة/مجموعة', 'set_ch_add', style='success'),
        btn('🔙 رجوع', 'admin_settings'))
  show(call, '📡 القنوات والمجموعات (يتم توليد رابط لكل واحدة عند القبول):', m)


def admins_view(call):
  rows = q_all('SELECT telegram_id, name FROM admins')
  lines = ['👥 الأدمنية:\n', '⭐ أساسيون (من متغيرات البيئة):']
  lines += [f'• {a}' for a in sorted(SUPER_ADMIN_IDS)]
  m = types.InlineKeyboardMarkup(row_width=1)
  if rows:
    lines.append('\n➕ مضافون من البوت:')
    for tid, name in rows:
      lines.append(f'• {name or ""} ({tid})')
      if is_super_admin(call.from_user.id):
        m.add(btn(f'🗑 حذف {name or tid}', f'set_adm_del_{tid}',
                  style='danger'))
  if is_super_admin(call.from_user.id):
    m.add(btn('➕ إضافة أدمن', 'set_adm_add', style='success'))
  m.add(btn('🔙 رجوع', 'admin_settings'))
  show(call, '\n'.join(lines), m)


def stickers_view(call):
  slots = [('welcome', '👋 ملصق الترحيب'), ('submitted', '🚀 ملصق استلام الطلب'),
           ('approved', '🎉 ملصق القبول'), ('rejected', '😢 ملصق الرفض')]
  m = types.InlineKeyboardMarkup(row_width=1)
  for slot, label in slots:
    on = '✅' if get_setting(f'sticker_{slot}') else '➖'
    m.add(btn(f'{label} {on}', f'set_stk_{slot}'))
  m.add(btn('🗑 حذف كل الملصقات', 'set_stk_clear', style='danger'),
        btn('🔙 رجوع', 'admin_settings'))
  show(call, '🎭 اختر خانة ثم أرسل لي الملصق اللي بدك ياه:', m)


def premium_view(call):
  mp, _ = pe_data()
  on = get_setting('premium_mode', 'on') == 'on'
  lines = ['🌟 الإيموجي المميز (Premium)\n']
  lines.append(f'الحالة: {"✅ مفعّل" if on else "⛔ موقّف"}')
  lines.append(f'عدد الإيموجيات المربوطة: {len(mp)}')
  if mp:
    lines.append(' '.join(mp.keys()))
  if PREMIUM_BROKEN['on']:
    lines.append('\n⚠️ تيليجرام رفض الإيموجي المميز فتم إيقافه مؤقتاً. تأكد أن'
                 ' حساب صاحب البوت (اللي أنشأه من BotFather) عليه Premium، ثم'
                 ' اضغط "إعادة المحاولة".')
    lines.append(f'السبب: {PREMIUM_BROKEN["reason"]}')
  m = types.InlineKeyboardMarkup(row_width=1)
  m.add(
      btn('➕ إضافة/تحديث إيموجيات مميزة', 'set_pe_add', style='success'),
      btn('👁 معاينة', 'set_pe_prev'),
      btn(f'⚙️ التفعيل: {"✅" if on else "⛔"} (اضغط للتبديل)', 'set_pe_toggle'),
  )
  if PREMIUM_BROKEN['on']:
    m.add(btn('🔄 إعادة المحاولة', 'set_pe_retry', style='primary'))
  m.add(btn('🗑 حذف كل الربط', 'set_pe_clear', style='danger'),
        btn('🔙 رجوع', 'admin_settings'))
  show(call, '\n'.join(lines), m)


def cb_settings(call, data):
  uid = call.from_user.id
  if data == 'set_pe':
    premium_view(call)
  elif data == 'set_pe_add':
    set_state(uid, 'awaiting_premium_emoji')
    m = types.InlineKeyboardMarkup().add(btn('✅ تم', 'set_pe_done',
                                              style='success'))
    send_msg(
        call.message.chat.id,
        '🌟 أرسل لي رسالة (أو عدة رسائل) فيها الإيموجيات المميزة اللي بدك'
        ' ياها.\nكل إيموجي مميز بيتربط تلقائياً بالإيموجي العادي المقابل له،'
        ' فمثلاً إيموجي 🎉 مميز بيحل محل كل 🎉 بالبوت.\n\n'
        f'📌 الإيموجيات المستخدمة بالبوت:\n{BOT_EMOJIS}\n\n'
        'لما تخلص اضغط تم (أو /cancel).',
        reply_markup=m,
    )
  elif data == 'set_pe_done':
    admin_input_states.pop(uid, None)
    premium_view(call)
  elif data == 'set_pe_prev':
    mp, _ = pe_data()
    if not mp:
      answer(call, 'ما في إيموجيات مربوطة بعد.', True)
      return
    send_msg(call.message.chat.id, 'معاينة: ' + ' '.join(mp.keys()))
  elif data == 'set_pe_toggle':
    set_setting('premium_mode',
                'off' if get_setting('premium_mode', 'on') == 'on' else 'on')
    premium_view(call)
  elif data == 'set_pe_retry':
    PREMIUM_BROKEN.update(on=False, reason='')
    premium_view(call)
  elif data == 'set_pe_clear':
    set_setting('pe_map', '')
    premium_view(call)
  elif data == 'admin_settings':

    show(call, settings_text(), settings_markup())
  elif data == 'set_branch':
    set_state(uid, 'awaiting_branch')
    send_msg(call.message.chat.id,
                     '🏷 اكتب اسم الفرع الجديد (أو /cancel):')
  elif data == 'set_owner':
    set_state(uid, 'awaiting_owner')
    send_msg(call.message.chat.id,
                     '👤 اكتب يوزر المسؤول بدون مسافات (أو /cancel):')
  elif data == 'set_channels':
    channels_view(call)
  elif data == 'set_ch_add':
    set_state(uid, 'awaiting_ch_add')
    send_msg(
        call.message.chat.id,
        '➕ حوّل لي أي رسالة من القناة/المجموعة، أو أرسل الـ ID تبعها.\n'
        '⚠️ لازم يكون البوت مشرف فيها مع صلاحية Invite Users.\n(أو /cancel)',
    )
  elif data.startswith('set_chren_'):
    set_state(uid, 'awaiting_ch_rename', id=int(data.rsplit('_', 1)[1]))
    send_msg(call.message.chat.id, '✏️ اكتب الاسم الجديد (أو /cancel):')
  elif data.startswith('set_chdely_'):
    cid = int(data.rsplit('_', 1)[1])
    q_exec('DELETE FROM channels WHERE id = ?', (cid,))
    audit(call.from_user, 'حذف قناة', cid)
    channels_view(call)
  elif data.startswith('set_chdel_'):
    cid = int(data.rsplit('_', 1)[1])
    m = types.InlineKeyboardMarkup(row_width=2)
    m.add(btn('🗑 نعم، احذف', f'set_chdely_{cid}', style='danger'),
          btn('↩️ تراجع', f'set_ch_{cid}'))
    show(call, '⚠️ متأكد من الحذف؟ (ما رح تنطرد أعضاء، بس ما رح تتولد روابطها'
         ' للطلاب الجدد)', m)
  elif data.startswith('set_ch_'):
    cid = int(data.rsplit('_', 1)[1])
    row = q_one('SELECT title, chat_id, kind FROM channels WHERE id = ?',
                (cid,))
    if not row:
      answer(call, '❌ غير موجودة.', True)
      return
    m = types.InlineKeyboardMarkup(row_width=2)
    m.add(btn('✏️ إعادة تسمية', f'set_chren_{cid}', style='primary'),
          btn('🗑 حذف', f'set_chdel_{cid}', style='danger'))
    m.add(btn('🔙 رجوع', 'set_channels'))
    show(call, f'{row[0]}\nID: {row[1]}\nالنوع: '
         f'{"قناة" if row[2] == "channel" else "مجموعة"}', m)
  elif data == 'set_admins':
    admins_view(call)
  elif data == 'set_adm_add':
    if not is_super_admin(uid):
      answer(call, '⚠️ للأدمن الأساسي فقط.', True)
      return
    set_state(uid, 'awaiting_admin_add')
    send_msg(call.message.chat.id,
                     '➕ أرسل Telegram ID للأدمن الجديد أو حوّل رسالة منه.')
  elif data.startswith('set_adm_del_'):
    if not is_super_admin(uid):
      answer(call, '⚠️ للأدمن الأساسي فقط.', True)
      return
    tid = int(data.rsplit('_', 1)[1])
    q_exec('DELETE FROM admins WHERE telegram_id = ?', (tid,))
    reload_admins()
    audit(call.from_user, 'حذف أدمن', tid)
    admins_view(call)
  elif data == 'set_autoxl':
    set_setting('auto_excel',
                'off' if get_setting('auto_excel', 'on') == 'on' else 'on')
    show(call, settings_text(), settings_markup())
  elif data == 'set_backup':
    set_setting('weekly_backup',
                'off' if get_setting('weekly_backup', 'on') == 'on' else 'on')
    show(call, settings_text(), settings_markup())
  elif data == 'set_stickers':
    stickers_view(call)
  elif data == 'set_stk_clear':
    for slot in ('welcome', 'submitted', 'approved', 'rejected'):
      set_setting(f'sticker_{slot}', '')
    stickers_view(call)
  elif data.startswith('set_stk_'):
    slot = data.rsplit('_', 1)[1]
    set_state(uid, 'awaiting_sticker', slot=slot)
    send_msg(call.message.chat.id, '🎭 أرسل الملصق الآن (أو /cancel):')
  answer(call)


@bot.callback_query_handler(func=lambda call: True)
def callback_router(call):
  data = call.data or ''
  uid = call.from_user.id
  try:
    # أزرار الطالب (ما بتحتاج صلاحية أدمن)
    if data == 'retry_freshman':
      return cb_retry(call)
    if data == 'mylinks':
      return cb_mylinks(call)

    if not is_admin(uid):
      answer(call, '⚠️ غير مصرح لك.', True)
      return

    if data == 'admin_panel':
      show(call, admin_panel_text(), get_admin_keyboard())
      return answer(call)
    if data == 'admin_stats':
      return cb_stats(call)
    if data == 'admin_archive':
      return cb_archive(call, 'all', 0)
    if data.startswith('admin_arch_'):
      _, _, status, page = data.split('_')
      return cb_archive(call, status, int(page))
    if data == 'admin_search':
      set_state(uid, 'awaiting_search')
      send_msg(call.message.chat.id,
                       '🔍 اكتب اسم الطالب أو رقم هاتفه أو Telegram ID (أو /cancel):')
      return answer(call)
    if data == 'admin_reset':
      set_state(uid, 'awaiting_reset')
      send_msg(call.message.chat.id,
                       '🔓 أرسل رقم هاتف الطالب أو Telegram ID لفك القفل عنه (أو /cancel):')
      return answer(call)
    if data in ('admin_export', 'admin_export_all'):
      answer(call, '⏳ جاري استخراج تقرير Excel...')
      send_excel_to([call.message.chat.id], data == 'admin_export')
      return
    if data == 'admin_zip':
      answer(call, '⏳ بدأ التجهيز...')
      Thread(target=export_photos_zip, args=(call.message.chat.id,),
             daemon=True).start()
      return
    if data == 'admin_broadcast':
      set_state(uid, 'awaiting_broadcast')
      send_msg(
          call.message.chat.id,
          '📣 أرسل الرسالة اللي بدك توصل لكل الطلاب المقبولين (نص/صورة/ملف)،'
          ' أو /cancel:')
      return answer(call)
    if data == 'bc_yes':
      return cb_broadcast_confirm(call)
    if data == 'bc_no':
      admin_input_states.pop(uid, None)
      edit_caption_or_text(call.message.chat.id, call.message.message_id,
                           '↩️ تم إلغاء الرسالة الجماعية.')
      return answer(call)
    if data == 'admin_log':
      return cb_log(call)
    if data == 'admin_settings' or data.startswith('set_'):
      return cb_settings(call, data)

    if data.startswith('view_file_'):
      send_student_file(call.message.chat.id, int(data.rsplit('_', 1)[1]))
      return answer(call)
    if data.startswith('approve_'):
      return cb_approve(call, int(data.split('_')[1]))
    if data.startswith('reject_menu_'):
      target = int(data.split('_')[2])
      edit_caption_or_text(
          call.message.chat.id, call.message.message_id,
          '📌 اختر سبب رفض الطلب ليتم إبلاغ الطالب تلقائياً:',
          reject_menu_markup(target),
      )
      return answer(call)
    if data.startswith('reject_reason_'):
      parts = data.split('_')
      return cb_reject_reason(call, int(parts[2]), parts[3])
    if data.startswith('resend_'):
      return cb_resend(call, int(data.split('_')[1]))
    if data.startswith('unban_'):
      return cb_ban(call, int(data.split('_')[1]), ban=False)
    if data.startswith('ban_'):
      return cb_ban(call, int(data.split('_')[1]), ban=True)
    answer(call)
  except Exception:
    logger.exception('callback error: %s', data)
    answer(call, '⚠️ صار خطأ غير متوقع، جرّب مرة ثانية.', True)


# ==================== أي رسالة ثانية من مستخدم عادي ====================
@bot.message_handler(
    content_types=ALL_CONTENT,
    func=lambda m: m.chat.type == 'private' and not is_admin(m.from_user.id),
)
def fallback(message):
  send_msg(message.chat.id, 'اكتب /start للبدء بالتوثيق 👋')


# ==================== التشغيل ====================
def run_bot():
  while True:
    try:
      bot.infinity_polling()
    except Exception:
      # لو الخيط طاح بصمت، البوت بيتوقف بينما Render شايف Flask شغالة.
      logger.exception('Polling stopped unexpectedly, restarting in 5s')
      time.sleep(5)


if __name__ == '__main__':
  logger.info('Freshmen bot started')
  Thread(target=run_bot, daemon=True).start()
  Thread(target=backup_loop, daemon=True).start()
  run_flask()
